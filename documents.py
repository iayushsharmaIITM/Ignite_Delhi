"""Document text extraction — one job: bytes in, plain text out.

WHY THIS IS ITS OWN FILE
The upload path is only as robust as this function, and this is the part that
touches untrusted input. Keeping it free of I/O and of FastAPI means it can be
tested directly with bytes, which is how the failure modes below were found.

DESIGN RULE: return errors as data, never raise them per file.
A batch upload where one file is a scanned PDF must still ingest the other nine.
So `extract()` raises for a single file and `extract_many()` converts that into a
per-file result. The caller decides what to do; nothing dies silently.

THE FAILURE MODE WORTH KNOWING ABOUT
A scanned (image-only) PDF extracts to an EMPTY STRING with no exception. pypdf
succeeds and hands back "". If that empty string is ingested, the brain silently
gains nothing and the user is told their upload worked. We check for it and
report it, because "your PDF is a scan, we cannot read it" is a real answer and
silence is not.

FORMATS
  .pdf            pypdf (page text; no OCR)
  .docx           python-docx (paragraphs + tables), imported optionally
  .txt .md .csv .json   decoded as text; Cognee's LLM does the structuring
"""

from __future__ import annotations

import io
import json
import os

# Per-file and per-upload ceilings. These exist to protect the demo: an
# unbounded upload would sit in the ingestion pipeline while a judge waits.
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_FILES = 20
MAX_TOTAL_CHARS = 500_000
MAX_FILE_CHARS = 200_000

TEXT_EXTS = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".rst"}

# Source and config files. These are plain text, and PS-2's Challenge paragraph
# names "code" as one of the five artefact types a company brain should connect.
# Without these the only honest answer to "where is the code?" was "we cannot
# read it" - which is exactly the gap the upload path was supposed to close.
CODE_EXTS = {
    ".py", ".js", ".ts", ".tsx", ".jsx", ".go", ".rs", ".java", ".rb", ".sh",
    ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".env.example",
    ".sql", ".tf", ".hcl", ".dockerfile", ".makefile", ".gradle",
}

SUPPORTED_EXTS = TEXT_EXTS | CODE_EXTS | {".pdf", ".docx"}

# Formats people will try that we deliberately do not support, with the reason.
# A specific message is worth more than a generic "unsupported".
KNOWN_UNSUPPORTED = {
    ".doc": "legacy binary .doc is not supported - re-save it as .docx or PDF",
    ".xls": "legacy .xls is not supported - export to .csv",
    ".xlsx": "export the sheet to .csv first",
    ".ppt": "export to PDF first",
    ".pptx": "export to PDF first",
    ".png": "images are not read - this build has no OCR",
    ".jpg": "images are not read - this build has no OCR",
    ".jpeg": "images are not read - this build has no OCR",
    ".zip": "extract the archive and upload the files individually",
}


class ExtractError(Exception):
    """A single file could not be turned into text. Message is user-facing."""


def extension(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def _decode(data: bytes) -> str:
    """Decode bytes that are claimed to be text.

    Tried in order because real uploads are messy: UTF-8 with a BOM is common
    from Windows Excel exports, and a stray latin-1 byte should degrade to a
    readable-ish string rather than reject the file outright.
    """
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ExtractError("PDF support is not installed (pip install pypdf)") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - pypdf raises many types
        raise ExtractError(f"could not open PDF: {str(exc)[:120]}") from exc

    pages = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001 - one bad page must not lose the rest
            text = ""
            _ = exc
        if text.strip():
            pages.append(f"[page {number}]\n{text.strip()}")

    if not pages:
        # The silent-empty case. Say what is actually wrong.
        raise ExtractError(
            "no extractable text - this PDF is probably a scan or image-only "
            "(no OCR in this build)"
        )
    return "\n\n".join(pages)


def _from_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise ExtractError(
            "DOCX support is not installed (pip install python-docx)"
        ) from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractError(f"could not open DOCX: {str(exc)[:120]}") from exc

    parts = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]

    # Tables carry real content in contracts and specs, and are easy to lose.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    if not parts:
        raise ExtractError("DOCX contained no text")
    return "\n".join(parts)


def extract(filename: str, data: bytes) -> str:
    """Turn one uploaded file into text. Raises ExtractError with a clear reason."""
    if not data:
        raise ExtractError("file is empty")

    if len(data) > MAX_FILE_BYTES:
        raise ExtractError(
            f"file is {len(data) / 1_048_576:.1f} MB - limit is "
            f"{MAX_FILE_BYTES // 1_048_576} MB"
        )

    ext = extension(filename)
    if not ext:
        raise ExtractError("file has no extension, so its type is unknown")
    if ext in KNOWN_UNSUPPORTED:
        raise ExtractError(KNOWN_UNSUPPORTED[ext])
    if ext not in SUPPORTED_EXTS:
        supported = ", ".join(sorted(SUPPORTED_EXTS))
        raise ExtractError(f"{ext} is not supported (supported: {supported})")

    if ext == ".pdf":
        text = _from_pdf(data)
    elif ext == ".docx":
        text = _from_docx(data)
    else:
        text = _decode(data)

    text = text.strip()
    if not text:
        raise ExtractError("no readable text found")

    if ext == ".json":
        # Pretty-print so the LLM sees structure rather than one long line.
        # A malformed JSON file is still worth ingesting as raw text, so this
        # is best-effort and never fatal.
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except Exception:  # noqa: BLE001
            pass

    if len(text) > MAX_FILE_CHARS:
        raise ExtractError(
            f"extracted text is {len(text):,} chars - limit is {MAX_FILE_CHARS:,}"
        )

    return text


def extract_many(files: list[tuple[str, bytes]]) -> tuple[list[dict], list[dict]]:
    """Extract a batch. Returns (documents, failures) - never raises.

    `documents` items: {"name", "text", "chars"}
    `failures`  items: {"name", "error"}
    """
    documents: list[dict] = []
    failures: list[dict] = []
    total_chars = 0

    for name, data in files[:MAX_FILES]:
        try:
            text = extract(name, data)
        except ExtractError as exc:
            failures.append({"name": name, "error": str(exc)})
            continue

        if total_chars + len(text) > MAX_TOTAL_CHARS:
            failures.append(
                {
                    "name": name,
                    "error": f"upload limit of {MAX_TOTAL_CHARS:,} characters reached",
                }
            )
            continue

        total_chars += len(text)
        documents.append({"name": name, "text": text, "chars": len(text)})

    for name, _ in files[MAX_FILES:]:
        failures.append(
            {"name": name, "error": f"only the first {MAX_FILES} files are ingested"}
        )

    return documents, failures
